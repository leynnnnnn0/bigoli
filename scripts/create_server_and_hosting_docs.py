from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/nathanielalvarez/Herd/bigoli")
SERVER_OUTPUT = ROOT / "StampBayan_Server_Requirements.docx"
HOSTING_OUTPUT = ROOT / "StampBayan_Hosting_Elsewhere_Confirmation.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
LIGHT_FILL = "F2F4F7"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D5DD")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_run(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or RGBColor(0, 0, 0)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def add_title(doc, title, subtitle=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    format_run(run, size=18, bold=True, color=BLUE)
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(12)
        run = sub.add_run(subtitle)
        format_run(run, size=11, color=RGBColor(80, 80, 80))


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    format_run(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else DARK_BLUE)


def add_body(doc, text, bold_label=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_label:
        run = paragraph.add_run(bold_label)
        format_run(run, bold=True)
        run = paragraph.add_run(text)
        format_run(run)
    else:
        run = paragraph.add_run(text)
        format_run(run)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.167
        run = paragraph.add_run("• ")
        format_run(run)
        run = paragraph.add_run(item)
        format_run(run)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.allow_autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        set_cell_border(cell)
        set_cell_shading(cell, LIGHT_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        format_run(run, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.10
            run = paragraph.add_run(value)
            format_run(run)
    return table


def add_key_value_table(doc, rows):
    return add_table(
        doc,
        ["Item", "Specification"],
        rows,
        [Inches(2.05), Inches(4.35)],
    )


def build_server_doc():
    doc = setup_doc()
    add_title(doc, "STAMPBAYAN SERVER REQUIREMENTS", "Website Platform Technical Hosting Specifications")

    add_key_value_table(
        doc,
        [
            ("Project", "StampBayan Website Platform for Bigoli"),
            ("Application Type", "Laravel 12 web application with Inertia.js, React, Vite, and MySQL/MariaDB database"),
            ("Purpose", "Define the minimum and recommended server specifications required to host, operate, secure, back up, and maintain the website platform."),
            ("Prepared For", "IT Review Team"),
            ("Prepared By", "Development Team"),
            ("Date", "June 18, 2026"),
        ],
    )

    add_heading(doc, "1. Application Stack")
    add_table(
        doc,
        ["Component", "Requirement"],
        [
            ("Backend", "PHP 8.2 or higher; Laravel 12 framework; Composer 2.x for dependency installation."),
            ("Frontend", "React 19, Inertia.js, TypeScript, Tailwind CSS, and Vite. Node.js is required for asset building during deployment, not for normal PHP request handling."),
            ("Database", "MySQL 8.0 or MariaDB 10.6 or higher recommended for production. SQLite is only suitable for local development/testing."),
            ("Web Server", "Nginx or Apache with the document root pointed to the Laravel public directory."),
            ("Storage", "Writable Laravel storage and bootstrap/cache directories; public storage link enabled for uploaded or generated files."),
            ("Background Jobs", "Database queue connection is supported. Production should run a queue worker or scheduled queue processing if queued tasks are enabled."),
        ],
        [Inches(1.7), Inches(4.7)],
    )

    add_heading(doc, "2. Minimum Server Specifications")
    add_table(
        doc,
        ["Resource", "Minimum"],
        [
            ("CPU", "2 vCPU or shared-hosting equivalent with stable PHP process capacity."),
            ("Memory", "2 GB RAM minimum for small production use. Shared hosting should provide enough PHP memory for Composer-built Laravel applications."),
            ("Storage", "10 GB SSD minimum, excluding long-term backups. More storage is needed if many QR/poster/PDF/image files are retained."),
            ("PHP Memory Limit", "256 MB minimum; 512 MB recommended for PDF, Excel, image, or large report generation."),
            ("Upload Limit", "20 MB minimum; adjust higher if business logos, QR poster assets, or exports require it."),
            ("Database", "1 production database with regular backups and remote or panel-based restore access."),
            ("SSL", "Valid HTTPS certificate for bigoli.stampbayan.com and related routes."),
        ],
        [Inches(1.7), Inches(4.7)],
    )

    add_heading(doc, "3. Recommended Production Specifications")
    add_table(
        doc,
        ["Resource", "Recommended"],
        [
            ("CPU", "2 to 4 vCPU for smoother dashboard, QR, PDF/export, and concurrent customer scan activity."),
            ("Memory", "4 GB RAM or higher for comfortable Laravel operation, deployment builds, and background tasks."),
            ("Storage", "25 GB SSD or higher with separate daily backup retention."),
            ("Database", "MySQL 8.0/MariaDB 10.6+ with automated daily backups, point-in-time or recent restore capability, and restricted user privileges."),
            ("Web Server", "Nginx with PHP-FPM preferred, or Apache with PHP-FPM/mod_php if configured securely."),
            ("Cache/Queue", "Database cache/session/queue are acceptable for the current application. Redis can be added later if traffic increases."),
            ("Email", "Authenticated SMTP provider for password reset, account notices, and future notification emails."),
        ],
        [Inches(1.7), Inches(4.7)],
    )

    add_heading(doc, "4. Required PHP Extensions and Services")
    add_bullets(
        doc,
        [
            "PHP extensions: BCMath, Ctype, cURL, DOM, Fileinfo, Filter, GD or Imagick, Intl, Mbstring, OpenSSL, PDO, PDO MySQL, Session, Tokenizer, XML, and Zip.",
            "Composer 2.x available during deployment or on the build machine.",
            "Node.js 20 LTS or 22 LTS and npm available during deployment/build, or built frontend assets uploaded from a trusted build machine.",
            "Cron scheduler access for Laravel scheduled tasks, using a once-per-minute cron entry if scheduled jobs are enabled.",
            "Queue worker support through supervisor/systemd, hosting panel process manager, or controlled cron fallback if queue tasks are used.",
            "Optional but recommended for PDF rendering: Chromium/Chrome support if server-side browser-based PDF generation is required.",
        ],
    )

    add_heading(doc, "5. Security Requirements")
    add_bullets(
        doc,
        [
            "HTTPS must be enforced on all login, dashboard, staff, customer, QR, and API routes.",
            "Production environment must use APP_ENV=production, APP_DEBUG=false, and a securely generated APP_KEY.",
            "Database credentials must be unique to this application and must not use root-level access.",
            "File and folder permissions must allow Laravel to write only to required directories such as storage and bootstrap/cache.",
            "Admin/hosting access should be limited to authorized maintainers and protected with strong passwords and two-factor authentication when available.",
            "Web application firewall, malware scanning, rate limiting, and regular security updates are recommended.",
        ],
    )

    add_heading(doc, "6. Backup and Recovery Requirements")
    add_table(
        doc,
        ["Area", "Requirement"],
        [
            ("Database Backups", "Automated daily backups with at least 7 to 30 days retention, depending on business policy."),
            ("File Backups", "Back up uploaded/generated files in storage/app/public and any business media assets."),
            ("Code Backups", "Maintain source code in a version-controlled repository and deploy from controlled releases."),
            ("Restore Testing", "IT should periodically test restoration of database and files to confirm backups are usable."),
            ("Incident Recovery", "Hosting provider should support rollback, backup restore, and access to logs for troubleshooting."),
        ],
        [Inches(1.7), Inches(4.7)],
    )

    add_heading(doc, "7. Deployment Requirements")
    add_bullets(
        doc,
        [
            "Point the domain/subdomain document root to the Laravel public directory.",
            "Configure the production .env file with correct APP_URL, database, mail, session, cache, queue, and filesystem settings.",
            "Run composer install --no-dev --optimize-autoloader during deployment.",
            "Run npm install and npm run build on the server or build assets locally and upload the generated public/build files.",
            "Run php artisan migrate --force after database backup and deployment approval.",
            "Run php artisan config:cache, route:cache, view:cache, and storage:link as part of production setup.",
        ],
    )

    add_heading(doc, "8. Current Hosting Suitability")
    add_body(
        doc,
        "The current Hostinger Premium Shared Hosting plan may be acceptable for an early production launch with approximately 500 customers if usage is light, mostly CRUD-based, and customer QR scans are not happening in very large simultaneous bursts. It should be monitored closely for response time, CPU throttling, memory limits, database performance, and queue/scheduler limitations.",
    )
    add_body(
        doc,
        "For better long-term sustainability, a VPS or cloud server with dedicated CPU/RAM, controlled queue workers, automated backups, and stronger monitoring is recommended once traffic, generated files, exports, or concurrent scan activity increases.",
    )

    doc.save(SERVER_OUTPUT)
    return SERVER_OUTPUT


def build_hosting_doc():
    doc = setup_doc()
    add_title(doc, "STAMPBAYAN HOSTING CONFIRMATION", "Hosting Portability and Developer Access Statement")

    add_key_value_table(
        doc,
        [
            ("Project", "StampBayan Website Platform for Bigoli"),
            ("Purpose", "Confirm whether the website can be hosted with another hosting provider while the current developer remains responsible for development and maintenance."),
            ("Prepared For", "IT Review Team"),
            ("Prepared By", "Development Team"),
            ("Date", "June 18, 2026"),
        ],
    )

    add_heading(doc, "Confirmation")
    add_body(
        doc,
        "Yes. The StampBayan website can be hosted with another hosting provider while the current developer remains in charge of development and maintenance, provided that the new hosting environment supports the application requirements and the developer is given the necessary access to deploy, configure, monitor, and maintain the system.",
    )

    add_heading(doc, "Required Access")
    add_bullets(
        doc,
        [
            "Hosting control panel or server access needed to manage PHP, database, files, SSL, domain/subdomain settings, cron jobs, and logs.",
            "Database access for migration, backup, restore, and troubleshooting.",
            "File manager, SFTP, SSH, Git deployment, or equivalent deployment access.",
            "Ability to configure environment variables, storage permissions, queue/scheduler setup, and production cache commands.",
            "Access to email/SMTP settings if the website will send password reset or notification emails.",
        ],
    )

    add_heading(doc, "Condition")
    add_body(
        doc,
        "The transfer should be scheduled and coordinated to avoid downtime. Before switching DNS, the new server should be prepared, tested, backed up, and verified using the production domain or a staging subdomain.",
    )

    doc.save(HOSTING_OUTPUT)
    return HOSTING_OUTPUT


if __name__ == "__main__":
    print(build_server_doc())
    print(build_hosting_doc())
