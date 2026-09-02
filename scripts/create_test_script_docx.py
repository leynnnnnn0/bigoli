from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/nathanielalvarez/Herd/bigoli/StampBayan_Test_Script.docx")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in edge_data.items():
                element.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_run(run, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def add_paragraph(doc, text="", bold_label=None, align=None, space_after=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    if align:
        paragraph.alignment = align

    if bold_label:
        label_run = paragraph.add_run(bold_label)
        format_run(label_run, bold=True)
        text_run = paragraph.add_run(text)
        format_run(text_run)
    else:
        run = paragraph.add_run(text)
        format_run(run)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    format_run(run, bold=True)
    return paragraph


def set_table_borders(table):
    border = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top=border,
                left=border,
                bottom=border,
                right=border,
            )
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def write_cell(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    format_run(run, bold=bold)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [Inches(1.8), Inches(4.7)])
    write_cell(table.rows[0].cells[0], rows[0][0], bold=True)
    write_cell(table.rows[0].cells[1], rows[0][1])
    for label, value in rows[1:]:
        cells = table.add_row().cells
        write_cell(cells[0], label, bold=True)
        write_cell(cells[1], value)
    set_table_borders(table)
    return table


def add_test_case_table(doc, cases):
    table = doc.add_table(rows=1, cols=6)
    headers = ["ID", "Module", "Scenario", "Test Steps", "Expected Result", "Result"]
    widths = [
        Inches(0.55),
        Inches(1.0),
        Inches(1.35),
        Inches(2.25),
        Inches(1.85),
        Inches(0.7),
    ]
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        write_cell(table.rows[0].cells[idx], header, bold=True)

    for case in cases:
        cells = table.add_row().cells
        for idx, value in enumerate(case):
            write_cell(cells[idx], value)

    set_table_borders(table)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("STAMPBAYAN WEBSITE PLATFORM TEST SCRIPT")
    format_run(run, bold=True)

    add_key_value_table(
        doc,
        [
            ("Project", "StampBayan Website Platform"),
            ("Document Type", "Test Script"),
            ("Purpose", "To verify core Bigoli StampBayan website functions, access controls, loyalty card workflows, stamp generation, stamp recording, card completion, and perk redemption before production use."),
            ("Prepared For", "IT Review Team"),
            ("Prepared By", "Development Team"),
            ("Version", "1.0"),
            ("Date", "June 18, 2026"),
        ],
    )

    add_heading(doc, "1. Test Scope")
    add_paragraph(
        doc,
        "This test script covers the main workflows of the Bigoli StampBayan website platform, including business account access, staff access, customer registration, loyalty card creation, stamp code generation, customer stamp recording, completed loyalty cards, perk claiming, and basic security validations.",
    )

    add_heading(doc, "2. Website Access Points")
    add_key_value_table(
        doc,
        [
            ("Main Website", "https://bigoli.stampbayan.com"),
            ("Business Login", "https://bigoli.stampbayan.com/login"),
            ("Customer Login", "https://bigoli.stampbayan.com/customer/login"),
            ("Staff Login", "https://bigoli.stampbayan.com/staff/login"),
        ],
    )

    add_heading(doc, "3. Test Login Credentials")
    add_key_value_table(
        doc,
        [
            ("Business Account", "Email: bigoli@gmail.com    Password: password"),
            ("Staff Account", "Username: Batangas    Password: password"),
            ("Customer Account", "Email: jashreil@gmail.com    Password: password"),
        ],
    )

    add_heading(doc, "4. Preconditions")
    preconditions = [
        "The Bigoli StampBayan website is accessible through the listed access points.",
        "A valid business account is available for testing.",
        "At least one Bigoli branch is available or can be created during testing.",
        "At least one staff account is available or can be created during testing.",
        "At least one customer account is available or can be created during testing.",
        "Test users have permission to create temporary loyalty cards, generate stamp codes, record stamps, and redeem perks.",
    ]
    for item in preconditions:
        add_paragraph(doc, item, bold_label="□ ")

    add_heading(doc, "5. Test Cases")
    cases = [
        ("TC-01", "Business Auth", "Business login", "Open https://bigoli.stampbayan.com/login; enter valid business email and password; submit.", "Business user is redirected to the business dashboard without error.", ""),
        ("TC-02", "Business Auth", "Invalid business login", "Open business login page; enter valid email with invalid password; submit.", "System rejects login and displays validation error.", ""),
        ("TC-03", "Business Dashboard", "Dashboard loading", "Login as business; open dashboard; review summary cards, customer counts, stamp activity, and branch filter if available.", "Dashboard loads successfully and displays business statistics.", ""),
        ("TC-04", "Branches", "Create branch", "Go to Branches; add Bigoli branch name/details; save.", "New branch appears in the branch list and is linked to the Bigoli business account.", ""),
        ("TC-05", "Branches", "Edit branch", "Open an existing branch; update branch details; save.", "Branch details are updated and remain linked to the same business.", ""),
        ("TC-06", "Staff", "Create staff account", "Go to Staff Accounts; select owned branch; enter staff username, email, and password; save.", "Staff account is created and appears in the staff account list.", ""),
        ("TC-07", "Staff", "Update staff status", "Open existing staff account; update active/inactive status; save.", "Staff status updates correctly and affects staff login permission.", ""),
        ("TC-08", "Staff Auth", "Staff login", "Open https://bigoli.stampbayan.com/staff/login; enter active staff credentials; submit.", "Staff dashboard opens successfully.", ""),
        ("TC-09", "Staff Auth", "Inactive staff login", "Set staff account to inactive; attempt login using that staff account.", "System denies access for inactive staff.", ""),
        ("TC-10", "Staff Dashboard", "Staff dashboard loading", "Login as staff; review assigned branch, available loyalty cards, issue stamp form, recent stamp activity, and available perk claims.", "Staff dashboard loads and displays only records for the staff member's assigned branch.", ""),
        ("TC-11", "Staff Access", "Branch-scoped cards", "Login as staff assigned to one branch; open issue stamp card selector and available card list.", "Staff can only select loyalty cards assigned to the staff branch.", ""),
        ("TC-12", "Staff Issue Stamp", "Staff gives stamp code", "Staff opens dashboard; selects loyalty card; enters reference number or order number; clicks Generate Code.", "Unique stamp code and QR code are generated under the staff account and selected branch.", ""),
        ("TC-13", "Staff Issue Stamp", "Staff presents QR code", "After staff generates a stamp code, display the generated QR code to the customer for scanning.", "Customer can scan the QR code and the same generated code is recorded as used after successful stamp recording.", ""),
        ("TC-14", "Staff Issue Stamp", "Missing staff reference", "Staff attempts to generate a stamp code without entering a reference number.", "System prevents generation and displays required field validation.", ""),
        ("TC-15", "Staff Issue Stamp", "Invalid staff card selection", "Staff attempts to generate a stamp using a loyalty card not assigned to the staff branch, if attempted through direct URL or modified request.", "System blocks generation and does not create a stamp code for an unauthorized card.", ""),
        ("TC-16", "Staff Perks", "Staff views perk claims", "Staff opens dashboard or perk claim section; review available, redeemed, and pending claim details for assigned branch.", "Staff sees only eligible perk claims for the assigned branch.", ""),
        ("TC-17", "Staff Perks", "Staff redeems perk", "Staff selects a customer perk claim; marks it as redeemed with optional remarks.", "Perk status changes to redeemed and shows staff redeemer information.", ""),
        ("TC-18", "Loyalty Cards", "Create loyalty card", "Business opens Loyalty Cards; create card with name, heading, subheading, expiry date, mechanics, colors, stamp count, perks, and branch assignment; save.", "Loyalty card template is saved with configured design, perks, and branch assignment.", ""),
        ("TC-19", "Loyalty Cards", "Preview loyalty card", "Open the created loyalty card; review preview, stamp slots, reward labels, mechanics, footer, and expiry information.", "Preview displays the correct card design, stamp count, and reward details.", ""),
        ("TC-20", "Loyalty Cards", "Edit loyalty card", "Open existing loyalty card; update colors, mechanics, perks, stamp count, or branch assignment; save.", "Loyalty card updates successfully and changes are reflected in the preview.", ""),
        ("TC-21", "QR Studio", "Save QR design", "Open QR Studio; update heading, subheading, logo/background if applicable, and colors; save.", "QR design settings are saved and displayed correctly.", ""),
        ("TC-22", "QR Studio", "Download QR", "Open QR Studio; download the QR code poster or QR image.", "QR file downloads successfully and points users to the customer registration flow.", ""),
        ("TC-23", "Issue Stamp", "Business generates code", "Business opens Issue Stamp; selects loyalty card; enters reference number; clicks Generate Code.", "Unique stamp code and QR code are generated.", ""),
        ("TC-24", "Issue Stamp", "Missing reference number", "Attempt to generate a stamp code without entering a reference number.", "System prevents generation and displays required field validation.", ""),
        ("TC-25", "Stamp Codes", "View generated codes", "Business opens Stamp Codes page; search/filter by code, status, branch, card, assignment, and date.", "Generated stamp codes are visible with correct status, branch, card, and reference number.", ""),
        ("TC-26", "Customer Auth", "Customer registration", "Open customer registration flow from QR or registration link; enter username, email, password, and business/branch context; submit.", "Customer account is created under the Bigoli business and correct branch if applicable.", ""),
        ("TC-27", "Customer Auth", "Customer login", "Open https://bigoli.stampbayan.com/customer/login; enter valid customer credentials; submit.", "Customer dashboard opens and displays available loyalty card(s).", ""),
        ("TC-28", "Customer Dashboard", "View active card", "Customer opens dashboard; review active loyalty card, stamp count, mechanics, rewards, and profile section.", "Customer dashboard displays card details and current stamp progress correctly.", ""),
        ("TC-29", "Stamp Recording", "Manual code entry", "Customer enters a valid unused stamp code manually and submits.", "Stamp is recorded and the customer stamp count increases by one.", ""),
        ("TC-30", "Stamp Recording", "QR scan", "Customer opens scan option; scans a valid generated stamp QR code.", "Stamp is recorded successfully and the dashboard updates.", ""),
        ("TC-31", "Stamp Recording", "Duplicate stamp code", "Customer attempts to enter or scan the same stamp code again.", "System rejects the duplicate code and does not add another stamp.", ""),
        ("TC-32", "Stamp Recording", "Expired stamp code", "Customer attempts to use an expired unused stamp code.", "System rejects the expired code and does not add a stamp.", ""),
        ("TC-33", "Stamp Security", "Wrong business code", "Customer attempts to use a code generated by another business.", "System rejects the code and does not record a stamp.", ""),
        ("TC-34", "Perks", "Unlock reward perk", "Customer records enough stamps to reach a configured perk stamp number.", "Perk claim is created and appears in the customer and business/staff perk claim views.", ""),
        ("TC-35", "Perks", "Business redeems perk", "Business opens Perk Claims; selects customer claim; marks as redeemed with optional remarks.", "Perk status changes to redeemed with redeemed date/user shown.", ""),
        ("TC-36", "Perks", "Undo redemption", "Business or staff opens redeemed claim; clicks undo redemption.", "Perk returns to available/unredeemed status.", ""),
        ("TC-37", "Completion", "Complete loyalty card", "Customer records stamps until the loyalty card reaches the required stamp count.", "Completed loyalty card record is created and visible in completed card/history section.", ""),
        ("TC-38", "Completion", "New card cycle", "After completion, customer continues recording new valid stamps for the same loyalty card.", "System starts or continues the next loyalty card cycle correctly.", ""),
        ("TC-39", "Customer Profile", "Update customer username", "Customer opens profile settings; updates username; saves.", "Customer username updates successfully on the dashboard.", ""),
        ("TC-40", "Customer Profile", "Update customer password", "Customer opens profile settings; enters current password, new password, and confirmation; saves.", "Password updates successfully and customer can login using the new password.", ""),
        ("TC-41", "Access Control", "Business data isolation", "Attempt to access or modify another business branch, staff, loyalty card, stamp code, or perk claim URL.", "System blocks access with unauthorized response or safe redirect.", ""),
        ("TC-42", "Logout", "Business/staff/customer logout", "Login as each role; click logout; attempt to access protected dashboard again.", "User is logged out and protected dashboard redirects to the correct login page.", ""),
    ]
    add_test_case_table(doc, cases)

    add_heading(doc, "6. Acceptance Criteria")
    acceptance = [
        "All critical test cases pass without server errors.",
        "No unauthorized user can access another business, branch, loyalty card, stamp code, or perk claim.",
        "Generated stamp codes cannot be reused after successful recording.",
        "Expired, used, or wrong-business stamp codes are rejected.",
        "Business, staff, and customer dashboards load within an acceptable response time under normal usage.",
        "All validation errors are clear and do not expose sensitive system information.",
    ]
    for item in acceptance:
        add_paragraph(doc, item, bold_label="□ ")

    add_heading(doc, "7. Sign-Off")
    add_key_value_table(
        doc,
        [
            ("Tester Name", ""),
            ("Tester Signature", ""),
            ("Date Tested", ""),
            ("Overall Result", "□ Passed    □ Failed    □ Passed with Remarks"),
            ("Remarks", ""),
        ],
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
